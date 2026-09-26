#!/usr/bin/env python3
"""Build the Finch pitch as a 16:9 PowerPoint."""

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

PAPER = RGBColor(0xF5, 0xF3, 0xED)
SURFACE = RGBColor(0xEF, 0xED, 0xE5)
CARD = RGBColor(0xFB, 0xFA, 0xF6)
INK = RGBColor(0x11, 0x18, 0x27)
INK2 = RGBColor(0x4B, 0x55, 0x63)
RULE = RGBColor(0xD8, 0xD5, 0xCD)
CORAL = RGBColor(0xF0, 0x5A, 0x47)
DARK = RGBColor(0x0E, 0x11, 0x16)
CREAM = RGBColor(0xF5, 0xF3, 0xED)
TEAL = RGBColor(0x0B, 0x7A, 0x73)
POSITIVE = RGBColor(0x2E, 0x7D, 0x57)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED_ON_DARK = RGBColor(0xC8, 0xC4, 0xBA)
DIM_ON_DARK = RGBColor(0x9A, 0x96, 0x8C)

DISPLAY = "Arial"
BODY = "Arial"
MONO = "Courier New"

W = 13.333
H = 7.5
OUT = Path(__file__).resolve().parent / "Finch-pitch.pptx"


def _font(run, name, size, bold, color):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.italic = False
    r_pr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = r_pr.find(qn(tag))
        if el is None:
            el = etree.SubElement(r_pr, qn(tag))
        el.set("typeface", name)


def _box(slide, l, t, w, h, anchor=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    anchor_map = {
        MSO_ANCHOR.TOP: "t",
        MSO_ANCHOR.MIDDLE: "ctr",
        MSO_ANCHOR.BOTTOM: "b",
    }
    tf._txBody.bodyPr.set("anchor", anchor_map[anchor])
    return tf


def _p(tf, text, size, bold, color, font, align=PP_ALIGN.LEFT, space_before=0, space_after=0, new=False):
    p = tf.add_paragraph() if new else tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    run = p.add_run()
    run.text = text
    _font(run, font, size, bold, color)
    return p


def text(slide, l, t, w, h, content, size=16, bold=False, color=INK, font=BODY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, link=None):
    tf = _box(slide, l, t, w, h, anchor)
    p = _p(tf, content, size, bold, color, font, align)
    if link:
        p.runs[0].hyperlink.address = link
    return tf


def runs(slide, l, t, w, h, paragraphs, anchor=MSO_ANCHOR.TOP):
    """paragraphs: list of list of (text, size, bold, color, font)."""
    tf = _box(slide, l, t, w, h, anchor)
    for i, parts in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        for part in parts:
            value, size, bold, color, font = part[:5]
            run = p.add_run()
            run.text = value
            _font(run, font, size, bold, color)
            if len(part) > 5:
                p.space_before = Pt(part[5])
    return tf


def rect(slide, l, t, w, h, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    shape.shadow.inherit = False
    return shape


def bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def header(slide, label, dark=False):
    ink = CREAM if dark else INK
    mute = DIM_ON_DARK if dark else INK2
    rect(slide, 0.52, 0.36, 0.16, 0.16, CORAL)
    text(slide, 0.78, 0.30, 2.2, 0.30, "Finch", 16, True, ink, DISPLAY, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, 7.4, 0.30, 5.4, 0.30, label.upper(), 11, False, mute, MONO, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)


def kicker(slide, label, l, t, w=6):
    text(slide, l, t, w, 0.26, label.upper(), 11, True, CORAL, MONO)


def footer_rule(slide, n, total, dark=False):
    color = RGBColor(0x2A, 0x2E, 0x36) if dark else RULE
    rect(slide, 0, 7.42, W * (n / total), 0.08, CORAL)
    rect(slide, W * (n / total), 7.42, W * (1 - n / total), 0.08, color)


def new_slide(prs, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide, DARK if dark else PAPER)
    return slide


def build():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.title = "Finch — Pitch"
    prs.core_properties.subject = "Transaction data intelligence"
    prs.core_properties.author = "Finch"
    total = 13

    # 1 Cover
    s = new_slide(prs)
    header(s, "Company pitch  ·  Sydney")
    kicker(s, "Transaction data intelligence", 0.55, 1.45)
    runs(
        s, 0.55, 1.85, 11.5, 2.5,
        [
            [("Raw transactions,", 40, True, INK, DISPLAY)],
            [("read like a merchant.", 40, True, CORAL, DISPLAY)],
            [("Understood like a customer.", 40, True, INK, DISPLAY)],
        ],
    )
    text(
        s, 0.55, 4.55, 8.6, 0.9,
        "Finch turns cryptic bank feed strings into clean merchant identity, category, location and behaviour — enriching every transaction for Australia’s banks, fintechs and retailers in real time.",
        16, False, INK2, BODY,
    )
    rect(s, 0.55, 5.7, 7.4, 0.015, RULE)
    metrics = [("4.2B+", "Transactions enriched"), ("95%", "Merchant accuracy"), ("42ms", "P95 latency")]
    for i, (num, label) in enumerate(metrics):
        x = 0.55 + i * 2.45
        text(s, x, 5.9, 2.2, 0.48, num, 26, True, INK, DISPLAY)
        text(s, x, 6.4, 2.2, 0.28, label.upper(), 10, False, INK2, MONO)
    footer_rule(s, 1, total)

    # 2 Story
    s = new_slide(prs)
    header(s, "The story")
    kicker(s, "100 words", 0.55, 0.85)
    runs(
        s, 0.55, 1.3, 12.2, 5.2,
        [
            [("Australian bank feeds still speak in codes. ", 22, False, INK, DISPLAY),
             ("SQ *GRN MRKT SYD", 18, True, CORAL, MONO),
             (" is a grocer in Sydney, but the line does not say so. Customers cannot explain their money. Credit models inherit the noise. Retailers cannot see who was paid.", 22, False, INK, DISPLAY)],
            [("Finch, in Sydney, reads every transaction the way a person would. In under fifty milliseconds a raw string becomes a merchant, a category, a place and a behaviour. One API then powers wellness, affordability, rewards and a live view of the market.", 22, False, INK, DISPLAY, 16)],
            [("The same primitives. Human review beside the models. 4.2 billion transactions already enriched. The feed, finally, makes sense.", 22, False, INK, DISPLAY, 16)],
        ],
    )
    text(s, 0.55, 6.85, 3, 0.3, "99 WORDS", 12, False, INK2, MONO)
    footer_rule(s, 2, total)

    # 3 Problem
    s = new_slide(prs)
    header(s, "The problem")
    kicker(s, "01", 0.55, 0.82)
    text(s, 0.55, 1.1, 12, 0.9, "Bank feeds were written for switches, not people.", 30, True, INK, DISPLAY)
    rect(s, 0.55, 2.2, 5.85, 4.85, DARK)
    text(s, 0.78, 2.38, 3.4, 0.28, "RAW  ·  BANK FEED", 11, False, DIM_ON_DARK, MONO)
    text(s, 4.7, 2.38, 1.4, 0.28, "AUD", 11, False, DIM_ON_DARK, MONO, PP_ALIGN.RIGHT)
    feeds = [
        ("SQ *GRN MRKT SYD 0426", "card"),
        ("UBER *TRIP HELP.UBER.COM", "card"),
        ("TFR AMZN AU 44118X", "transfer"),
        ("POS SUSHI HUB CBD", "point of sale"),
        ("WWW NETFLIX.COM 899", "recurring"),
        ("DD ORIGIN ENERGY", "direct debit"),
    ]
    for i, (raw, kind) in enumerate(feeds):
        y = 2.82 + i * 0.64
        rect(s, 0.78, y, 5.4, 0.01, RGBColor(0x2A, 0x2E, 0x36))
        text(s, 0.78, y + 0.08, 3.7, 0.36, raw, 12, False, CREAM, MONO, anchor=MSO_ANCHOR.MIDDLE)
        text(s, 4.4, y + 0.08, 1.7, 0.36, kind, 11, False, DIM_ON_DARK, MONO, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)
    pains = [
        ("01", "Customers can’t see their own money.", "A truncated merchant string is not a brand, a category, or a place. Wellness features stall on lines nobody can explain."),
        ("02", "Credit decisions inherit the noise.", "Income, bills, subscriptions and affordability are in the history — locked inside descriptions built for settlement, not underwriting."),
        ("03", "The market stays invisible.", "Without a clean merchant, retailers and rewards teams cannot see share of wallet, cohorts, or where spend is actually moving."),
    ]
    for i, (num, title, body) in enumerate(pains):
        y = 2.2 + i * 1.62
        rect(s, 6.7, y, 6.05, 0.015, INK)
        text(s, 6.7, y + 0.1, 1.2, 0.24, num, 11, False, CORAL, MONO)
        text(s, 6.7, y + 0.36, 6.0, 0.4, title, 16, True, INK, DISPLAY)
        text(s, 6.7, y + 0.8, 6.0, 0.7, body, 13, False, INK2, BODY)
    footer_rule(s, 3, total)

    # 4 Enrichment
    s = new_slide(prs)
    header(s, "Transaction intelligence engine  ·  Live")
    kicker(s, "02  —  Enrichment", 0.55, 0.8)
    text(s, 0.55, 1.08, 7.2, 1.05, "The same line, read the way a customer would.", 28, True, INK, DISPLAY)
    text(s, 8.0, 1.15, 4.8, 1.0, "Clean merchant identity, category, location and behaviour attached to every card, direct debit and transfer description — including truncated Australian feeds.", 13, False, INK2, BODY)
    rows = [
        ("SQ *GRN MRKT SYD 0426", "Green Market", "Groceries  ·  Sydney  ·  Weekly recurring", "96%"),
        ("UBER *TRIP HELP.UBER.COM", "Uber", "Transport  ·  Melbourne  ·  Recurring behaviour", "98%"),
        ("TFR AMZN AU 44118X", "Amazon AU", "Retail  ·  Online  ·  Subscription detected", "94%"),
        ("POS SUSHI HUB CBD", "Sushi Hub", "Dining  ·  Sydney CBD  ·  New merchant", "91%"),
    ]
    rect(s, 0.55, 2.35, 12.25, 0.015, INK)
    for i, (raw, name, meta, conf) in enumerate(rows):
        y = 2.45 + i * 1.0
        rect(s, 0.55, y + 0.9, 12.25, 0.012, RULE)
        rect(s, 0.55, y + 0.16, 3.55, 0.55, SURFACE)
        text(s, 0.68, y + 0.16, 3.3, 0.55, raw, 12, False, INK2, MONO, anchor=MSO_ANCHOR.MIDDLE)
        text(s, 4.4, y + 0.12, 0.4, 0.55, "→", 18, True, CORAL, DISPLAY, anchor=MSO_ANCHOR.MIDDLE)
        text(s, 4.9, y + 0.08, 5.6, 0.38, name, 18, True, INK, DISPLAY)
        text(s, 4.9, y + 0.46, 5.6, 0.32, meta, 13, False, INK2, BODY)
        rect(s, 11.15, y + 0.22, 1.15, 0.42, RGBColor(0xE5, 0xF2, 0xEA))
        text(s, 11.15, y + 0.22, 1.15, 0.42, conf, 13, True, POSITIVE, MONO, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    text(s, 0.55, 6.6, 12, 0.28, "WORKED EXAMPLE FROM THE FINCH PLATFORM   ·   4 / 4 MERCHANTS MATCHED   ·   94.8% AVG CONFIDENCE", 11, False, INK2, MONO)
    footer_rule(s, 4, total)

    # 5 Pipeline
    s = new_slide(prs, dark=True)
    header(s, "Real-time intelligence", dark=True)
    kicker(s, "03  —  Under fifty milliseconds", 0.55, 0.82)
    text(s, 0.55, 1.12, 7.4, 1.15, "One transaction. Five stages.", 32, True, CREAM, DISPLAY)
    text(s, 8.15, 1.2, 4.6, 1.05, "Every line through the Finch API is matched, categorised, read for behaviour, and written into aggregated market intelligence.", 14, False, MUTED_ON_DARK, BODY)
    stages = [
        ("T+0ms", "Raw feed", "The bank line arrives, still cryptic.", "SQ *GRN MRKT SYD 0426"),
        ("T+18ms", "Merchant match", "Identity resolved against the merchant graph.", "Green Market\ngreenmarket.com.au"),
        ("T+27ms", "Categorisation", "Category, cadence and place assigned.", "Groceries · Recurring · Weekly"),
        ("T+34ms", "Behaviour", "Signals pulled from the customer’s history.", "Weekly spend +8% vs 4wk avg"),
        ("T+42ms", "Intelligence", "The segment updates for the network.", "Health-conscious grocery segment"),
    ]
    for i, (when, title, body, ex) in enumerate(stages):
        x = 0.5 + i * 2.54
        rect(s, x, 2.6, 2.4, 4.35, RGBColor(0x17, 0x1B, 0x22), RGBColor(0x2A, 0x2E, 0x36))
        text(s, x + 0.14, 2.76, 2.12, 0.28, when, 12, False, CORAL, MONO)
        text(s, x + 0.14, 3.2, 2.12, 0.7, title, 16, True, CREAM, DISPLAY)
        text(s, x + 0.14, 3.95, 2.12, 1.15, body, 13, False, MUTED_ON_DARK, BODY)
        text(s, x + 0.14, 5.5, 2.12, 1.15, ex, 12, False, RGBColor(0x5E, 0xE0, 0xD0), MONO)
    footer_rule(s, 5, total, dark=True)

    # 6 Platform
    s = new_slide(prs)
    header(s, "POST /v1/transactions/enrich  ·  REST  ·  AU-hosted")
    kicker(s, "04  —  The platform", 0.55, 0.8)
    text(s, 0.55, 1.08, 7.0, 1.0, "One API. Three intelligence layers.", 30, True, INK, DISPLAY)
    text(s, 7.8, 1.15, 5.0, 0.95, "A single integration scales from a merchant lookup to a full behavioural layer. Each product uses the same primitives and can be adopted on its own.", 13, False, INK2, BODY)
    cards = [
        ("01  ·  ENRICHMENT API", "Turn a bank string into a merchant.", "Identity, category, location, logo and behaviour across Australian and global card, debit and transfer descriptions.", "Merchant name   Category   Location   Logo & website   Recurring   Confidence"),
        ("02  ·  ANALYTICS API", "Understand how a customer spends.", "Enriched history becomes income, bills, subscriptions, discretionary versus essential spend, and affordability.", "Income   Bills   Subscriptions   Fee saver   Cash-flow   Affordability"),
        ("03  ·  INTELLIGENCE API", "See markets and merchants move.", "Privacy-preserved signals across the Finch network — category, region, share of wallet and cohort behaviour as a feed.", "Spend   Retail share   Merchants   Benchmarks   Geography   Alerts"),
    ]
    for i, (idx, title, body, tags) in enumerate(cards):
        x = 0.5 + i * 4.22
        rect(s, x, 2.4, 4.05, 4.55, CARD, RULE)
        text(s, x + 0.22, 2.58, 3.6, 0.28, idx, 11, False, CORAL, MONO)
        text(s, x + 0.22, 3.0, 3.6, 1.05, title, 18, True, INK, DISPLAY)
        text(s, x + 0.22, 4.15, 3.6, 1.35, body, 13, False, INK2, BODY)
        text(s, x + 0.22, 5.7, 3.6, 0.95, tags, 12, False, INK, BODY)
    footer_rule(s, 6, total)

    # 7 Solutions
    s = new_slide(prs)
    header(s, "Solutions")
    kicker(s, "05", 0.55, 0.8)
    text(s, 0.55, 1.08, 7.6, 1.05, "Built for the problems Australian financial teams actually have.", 26, True, INK, DISPLAY)
    text(s, 8.3, 1.15, 4.5, 1.0, "Every solution sits on the same enrichment and analytics primitives — a money feature and a credit decision share one underlying dataset.", 13, False, INK2, BODY)
    sols = [
        ("01", "Financial wellness", "Customers can’t see or explain their own money. Finch returns cash-flow, bills and subscriptions.", "+27% engagement in customer money features."),
        ("02", "Affordability", "Underwriting still leans on outdated income snapshots. Finch verifies income and expenditure from live transactions.", "Faster credit decisions with lower default rates."),
        ("03", "Rewards", "Generic rewards miss real behaviour. Offers key off identified merchants and category patterns.", "Higher activation and repeat redemption."),
        ("04", "Retail intelligence", "Retail teams lack a live view of the Australian consumer. Benchmarks arrive by category, region and merchant.", "Sharper merchandising, pricing and expansion."),
    ]
    rect(s, 0.55, 2.35, 12.25, 0.015, INK)
    for i, (num, title, body, outcome) in enumerate(sols):
        y = 2.45 + i * 1.15
        rect(s, 0.55, y + 1.05, 12.25, 0.012, RULE)
        text(s, 0.55, y + 0.18, 0.6, 0.4, num, 13, False, CORAL, MONO)
        text(s, 1.3, y + 0.14, 3.3, 0.8, title, 18, True, INK, DISPLAY, anchor=MSO_ANCHOR.MIDDLE)
        text(s, 4.7, y + 0.12, 4.3, 0.85, body, 13, False, INK2, BODY, anchor=MSO_ANCHOR.MIDDLE)
        text(s, 9.15, y + 0.12, 3.6, 0.85, outcome, 14, True, TEAL, BODY, anchor=MSO_ANCHOR.MIDDLE)
    footer_rule(s, 7, total)

    # 8 Trust
    s = new_slide(prs)
    header(s, "Open banking  ·  Security")
    kicker(s, "06", 0.55, 0.82)
    text(s, 0.55, 1.12, 11.5, 1.0, "The open banking intelligence layer, held to a bank standard.", 28, True, INK, DISPLAY)
    text(s, 0.55, 2.35, 5.5, 0.3, "CONSUMER DATA RIGHT", 12, True, INK, MONO)
    text(s, 7.15, 2.35, 5.5, 0.3, "CONTROLS", 12, True, INK, MONO)
    left = [
        ("Accreditation", "CDR-aligned infrastructure"),
        ("Consent", "Fine-grained and revocable"),
        ("Recipients", "ADI, ADR and trusted partners"),
        ("Purpose", "Data minimisation, purpose-bound use"),
    ]
    right = [
        ("In transit", "TLS 1.3 on every endpoint"),
        ("At rest", "AES-256 for stored datasets"),
        ("Access", "Role-based permissions, immutable audit"),
        ("Residency", "Australian regions  ·  99.99% SLA"),
    ]
    for i, ((a, b), (c, d)) in enumerate(zip(left, right)):
        y = 2.8 + i * 0.95
        rect(s, 0.55, y, 5.9, 0.012, RULE)
        rect(s, 7.15, y, 5.6, 0.012, RULE)
        text(s, 0.55, y + 0.16, 2.1, 0.5, a, 15, True, INK, DISPLAY, anchor=MSO_ANCHOR.MIDDLE)
        text(s, 2.7, y + 0.16, 3.6, 0.5, b, 15, False, INK2, BODY, anchor=MSO_ANCHOR.MIDDLE)
        text(s, 7.15, y + 0.16, 1.8, 0.5, c, 15, True, INK, DISPLAY, anchor=MSO_ANCHOR.MIDDLE)
        text(s, 9.05, y + 0.16, 3.6, 0.5, d, 15, False, INK2, BODY, anchor=MSO_ANCHOR.MIDDLE)
    footer_rule(s, 8, total)

    # 9 Scale
    s = new_slide(prs, dark=True)
    header(s, "By the numbers", dark=True)
    kicker(s, "07", 0.55, 0.82)
    text(s, 0.55, 1.12, 10, 1.05, "Scale that powers Australia’s financial data infrastructure.", 28, True, CREAM, DISPLAY)
    nums = [
        ("4.2B+", "Transactions enriched to date", False),
        ("10M+", "End customers analysed", False),
        ("95%", "Enrichment accuracy", True),
        ("1.5M+", "Merchants in the database", False),
    ]
    for i, (num, label, coral) in enumerate(nums):
        col, row = i % 2, i // 2
        x = 0.55 + col * 6.2
        y = 2.45 + row * 1.55
        text(s, x, y, 5.5, 0.8, num, 48, True, CORAL if coral else CREAM, DISPLAY)
        text(s, x, y + 0.85, 5.5, 0.35, label.upper(), 12, False, DIM_ON_DARK, MONO)
    rect(s, 0.55, 5.65, 12.2, 0.012, RGBColor(0x2A, 0x2E, 0x36))
    text(s, 0.55, 5.85, 4, 0.7, "$2B+", 44, True, CREAM, DISPLAY)
    text(s, 5.0, 6.05, 7.5, 0.5, "IN TRANSACTION VALUE PROCESSED EVERY MONTH ACROSS AUSTRALIAN BANKS AND FINTECHS", 12, False, DIM_ON_DARK, MONO, anchor=MSO_ANCHOR.MIDDLE)
    footer_rule(s, 9, total, dark=True)

    # 10 Proof
    s = new_slide(prs)
    header(s, "Recognition")
    kicker(s, "08  —  What partners say", 0.55, 0.78)
    text(s, 0.55, 1.15, 12.2, 1.35, "“Finch improved our merchant identification accuracy overnight. Categorisation is a step-change beyond anything we’d built in-house.”", 22, True, INK, DISPLAY)
    rect(s, 0.55, 2.6, 8.2, 0.012, RULE)
    text(s, 0.55, 2.72, 4, 0.32, "Amelia Rowntree", 15, True, INK, DISPLAY)
    text(s, 5.2, 2.72, 4.5, 0.32, "Head of Product  ·  Meridian Bank", 14, False, INK2, BODY, PP_ALIGN.RIGHT)
    awards = [
        ("2019", "Money20/20 USA Startup Pitch", "FINALIST"),
        ("2019", "Most Innovative Fintech Team, Finder Awards", "WINNER"),
        ("2018", "Diversity in Fintech Rising Star, YBF Ventures", "WINNER"),
        ("2018", "Best Personal Finance App, Finder Awards", "WINNER"),
        ("2018", "Best Digital Wallet Australia, Finnies", "WINNER"),
    ]
    for i, (year, name, result) in enumerate(awards):
        x = 0.5 + i * 2.54
        rect(s, x, 3.3, 2.42, 2.15, CARD, RULE)
        text(s, x + 0.12, 3.42, 2.15, 0.26, year, 12, False, CORAL, MONO)
        text(s, x + 0.12, 3.75, 2.15, 1.05, name, 13, True, INK, BODY)
        text(s, x + 0.12, 4.95, 2.15, 0.28, result, 11, False, TEAL, MONO)
    text(s, 0.55, 5.7, 1.5, 0.28, "TRUSTED BY", 11, False, INK2, MONO, anchor=MSO_ANCHOR.MIDDLE)
    logos = "Meridian Bank    Payloft    AURA Fintech    Ledgerly    Northwind Retail    Cascade Pay    Beacon Credit    Terrace Capital    Coastal Union    Harbour Fintech"
    text(s, 2.15, 5.65, 10.6, 0.7, logos, 13, True, INK2, DISPLAY)
    footer_rule(s, 10, total)

    # 11 Why
    s = new_slide(prs)
    header(s, "Why Finch")
    kicker(s, "09", 0.55, 0.8)
    text(s, 0.55, 1.08, 7.4, 1.0, "Human intelligence and machine learning, one pipeline.", 26, True, INK, DISPLAY)
    text(s, 8.15, 1.12, 4.65, 1.0, "A Sydney data science team designs, monitors and improves the models that classify every transaction — and holds accuracy at enterprise levels.", 13, False, INK2, BODY)
    whys = [
        ("01", "An Australian merchant graph.", "1.5 million merchants, with name, category, location, logo and site — built for the way local card, Osko and direct-debit strings are actually written."),
        ("02", "Accuracy with a human loop.", "Models do the volume. People review the edge. That combination is what keeps enrichment at 95% instead of drifting quietly in production."),
        ("03", "Fast enough to sit in the product.", "42ms at the 95th percentile. Sandbox keys in minutes. Production with observability, a 99.99% SLA, and data that stays in Australia."),
        ("04", "One integration, four problems.", "Wellness, affordability, rewards and retail do not need four vendors. They adopt layers of the same API as the use case grows."),
    ]
    for i, (num, title, body) in enumerate(whys):
        col, row = i % 2, i // 2
        x = 0.5 + col * 6.4
        y = 2.4 + row * 2.3
        rect(s, x, y, 6.2, 2.15, CARD, RULE)
        text(s, x + 0.22, y + 0.16, 1.2, 0.26, num, 12, False, CORAL, MONO)
        text(s, x + 0.22, y + 0.48, 5.75, 0.45, title, 18, True, INK, DISPLAY)
        text(s, x + 0.22, y + 1.05, 5.75, 0.9, body, 13, False, INK2, BODY)
    footer_rule(s, 11, total)

    # 12 Start
    s = new_slide(prs)
    header(s, "How it works")
    kicker(s, "10", 0.55, 0.82)
    text(s, 0.55, 1.12, 12, 0.7, "From a raw file to a personalised experience.", 28, True, INK, DISPLAY)
    steps = [
        ("01", "Connect", "Integrate with the Finch API and stream transaction data securely."),
        ("02", "Enrich", "Every line is matched, categorised and tagged with merchant intelligence."),
        ("03", "Analyse", "Behaviour, income and affordability signals are computed continuously."),
        ("04", "Personalise", "Signals land in products, credit decisions and customer experiences."),
    ]
    for i, (num, title, body) in enumerate(steps):
        x = 0.55 + i * 3.2
        rect(s, x, 2.15, 2.95, 0.015, INK)
        text(s, x, 2.3, 2.9, 0.26, num, 12, False, CORAL, MONO)
        text(s, x, 2.62, 2.9, 0.4, title, 20, True, INK, DISPLAY)
        text(s, x, 3.15, 2.9, 1.1, body, 14, False, INK2, BODY)
    rect(s, 0.5, 4.6, 12.35, 2.4, DARK)
    text(s, 0.75, 4.82, 3.3, 0.4, "Start with your own file.", 16, True, CREAM, DISPLAY)
    text(s, 0.75, 5.28, 3.3, 0.8, "A free enrichment report, usually back within one business day.", 13, False, MUTED_ON_DARK, BODY)
    bits = [
        ("YOU SEND", "An anonymised CSV, or a synthetic sample."),
        ("YOU RECEIVE", "Per-line enrichment, category accuracy, recurring detection."),
        ("THEN", "A 30-minute walkthrough. Samples stay under NDA and are deleted after."),
    ]
    for i, (label, body) in enumerate(bits):
        x = 4.3 + i * 2.8
        text(s, x, 4.85, 2.6, 0.28, label, 11, False, CORAL, MONO)
        text(s, x, 5.25, 2.6, 1.3, body, 13, False, CREAM, BODY)
    footer_rule(s, 12, total)

    # 13 Close
    s = new_slide(prs, dark=True)
    header(s, "Finch Pty Ltd  ·  Sydney", dark=True)
    kicker(s, "Start with Finch", 0.55, 1.15)
    text(s, 0.55, 1.6, 11.5, 2.0, "Turn transaction data into a competitive advantage.", 36, True, CREAM, DISPLAY)
    text(s, 0.55, 3.8, 9, 0.8, "Book a walkthrough. We’ll show Finch running on real Australian transaction data and scope a path to production.", 16, False, MUTED_ON_DARK, BODY)
    rect(s, 0.55, 4.85, 2.35, 0.52, CORAL)
    text(s, 0.55, 4.85, 2.35, 0.52, "Request a demo", 14, True, WHITE, BODY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, link="https://finchxp.tech/request-demo")
    text(s, 3.1, 4.85, 3.2, 0.52, "hello@finchxp.com", 14, False, MUTED_ON_DARK, MONO, anchor=MSO_ANCHOR.MIDDLE, link="mailto:hello@finchxp.com")
    text(s, 6.4, 4.85, 3, 0.52, "finchxp.tech", 14, False, MUTED_ON_DARK, MONO, anchor=MSO_ANCHOR.MIDDLE, link="https://finchxp.tech/")
    rect(s, 0.55, 5.85, 12.2, 0.012, RGBColor(0x2A, 0x2E, 0x36))
    for i, (num, label) in enumerate(metrics):
        x = 0.55 + i * 3.3
        text(s, x, 6.1, 3, 0.42, num, 22, True, CREAM, DISPLAY)
        text(s, x, 6.55, 3, 0.28, label.upper(), 11, False, DIM_ON_DARK, MONO)
    footer_rule(s, 13, total, dark=True)

    prs.save(OUT)
    print(f"wrote {OUT} ({OUT.stat().st_size} bytes, {len(prs.slides)} slides)")


if __name__ == "__main__":
    build()
