#!/usr/bin/env python3
"""Build the sample Red Herring Prospectus .docx from labeled fixture data."""

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from sample_data import COMPANY, PEOPLE, TICKETS, gold_entities

OUT = Path("data/red_herring_prospectus.docx")


def add_centered(document, text, style=None, bold=False, size=None):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if style:
        paragraph.style = style
    run = paragraph.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    return paragraph


def build() -> Path:
    document = Document()
    section = document.sections[0]
    header = section.header.paragraphs[0]
    header.text = "CONFIDENTIAL — SAMPLE FILING FOR PII REDACTION"
    footer = section.footer.paragraphs[0]
    footer.text = f"{COMPANY['legal_name']} — Red Herring Prospectus"

    add_centered(document, "RED HERRING PROSPECTUS", bold=True, size=18)
    add_centered(
        document,
        f"Dated: {COMPANY['prospectus_date']}  |  Please read Section 32 of the Companies Act, 2013",
        size=10,
    )
    add_centered(document, "Book Built Issue", bold=True, size=12)
    document.add_paragraph()
    add_centered(document, COMPANY["legal_name"], bold=True, size=16)

    document.add_paragraph(
        f"Corporate Identity Number (CIN): {COMPANY['cin']}. Our Company was incorporated on "
        f"{COMPANY['incorporation_date']} under the Companies Act, 2013. This identifier, the "
        "prospectus date, and the incorporation date are corporate facts, not personal data."
    )
    document.add_paragraph(f"Registered Office: {COMPANY['registered_office']}")
    document.add_paragraph(f"Corporate Office: {COMPANY['corporate_office']}")
    document.add_paragraph(
        f"Tel: {COMPANY['phone']}    Email: {COMPANY['email']}    "
        f"Website: {COMPANY['website']}"
    )
    document.add_paragraph(
        f"Contact Person: {PEOPLE[0]['name']}, Company Secretary and Compliance Officer"
    )
    document.add_paragraph(
        f"OUR PROMOTERS: {PEOPLE[0]['name']} and {PEOPLE[1]['name']}"
    )

    document.add_heading("The Offer", level=1)
    document.add_paragraph(
        f"Initial public offering of equity shares of face value of {COMPANY['face_value']} each "
        f"aggregating up to {COMPANY['offer_size']} (the “Offer”). Application number APP-908812 "
        "is an operational identifier. Ticket and Order numbers in Appendix A are also operational "
        "IDs and are not treated as personally identifiable information."
    )
    document.add_paragraph(
        "SEBI, BSE and NSE are regulators / stock exchanges, not company names of the issuer. "
        "Nothing in this section should redact the words Offer, Ticket, Order, or Section 32."
    )

    document.add_heading("Board of Directors", level=1)
    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    headers = ["Name", "Designation", "Date of Birth", "Email", "Phone", "Address"]
    for i, title in enumerate(headers):
        table.rows[0].cells[i].text = title
    for person in PEOPLE:
        row = table.add_row().cells
        row[0].text = person["name"]
        row[1].text = person["role"]
        row[2].text = f"Date of Birth: {person['dob']}"
        row[3].text = person["email"]
        row[4].text = person["phone"]
        row[5].text = person["address"]

    document.add_paragraph(
        f"{PEOPLE[1]['name']} is a partner at {PEOPLE[1]['company']}. "
        f"{PEOPLE[3]['name']} previously founded {PEOPLE[3]['company']}."
    )

    document.add_heading("Appendix A — Investor Support Ticket Log", level=1)
    document.add_paragraph(
        "The following tickets were raised by directors / officers during KYC. "
        "Ticket IDs (TKT-*) and Order IDs (ORD-*) must remain visible after redaction."
    )

    for ticket in TICKETS:
        person = ticket["person"]
        document.add_heading(f"Ticket {ticket['id']}", level=2)
        document.add_paragraph(f"Ticket ID: {ticket['id']}    Order ID: {ticket['order_id']}    Status: Open")
        document.add_paragraph(f"From: {person['name']} <{person['email']}>")
        if person.get("ip"):
            document.add_paragraph(f"Source IP: {person['ip']}")
        document.add_paragraph(f"Phone: {person['phone']}")
        document.add_paragraph(f"Address: {person['address']}")

        extras = []
        if person.get("ssn"):
            extras.append(f"SSN {person['ssn']}")
        if person.get("cc"):
            extras.append(f"card {person['cc']}")
        extra_text = (" Please verify " + " and ".join(extras) + ".") if extras else ""
        document.add_paragraph(
            f"Hello team, this is {person['name']}. Kindly update KYC. "
            f"I was born on {person['dob']}.{extra_text} "
            f"Do not change Ticket {ticket['id']} or Order {ticket['order_id']}."
        )

    document.add_paragraph()
    document.add_paragraph(
        f"For and on behalf of {COMPANY['legal_name']}"
    )
    document.add_paragraph(f"Sd/-")
    document.add_paragraph(f"{PEOPLE[0]['name']}")
    document.add_paragraph("Company Secretary and Compliance Officer")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUT)
    Path("data/gold_labels.json").write_text(
        json.dumps(gold_entities(), indent=2),
        encoding="utf-8",
    )
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
